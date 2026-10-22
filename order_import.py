"""Импорт старой таблицы заказов (Excel «Выжимальня.xlsx») в CRM.

Формат листа «Заказы»: B — дата заказа, C — дата выдачи, D…S — масла (в ячейке объём
бутылки в мл; 2000 — это несколько бутылок), T — допы, U — ФИО, V — доставка/канал,
W — комментарий, X — телефон, Y — откуда. Строка без даты и имени — ещё бутылки того же
заказа; строки с формулами и нижний блок «общая/шт/Прибыль» — итоги, их пропускаем.

Цен в строках нет: сумма считается по листу «Прайс» той же книги (это оценка по текущему
прайсу), а если масла там нет — по прайс-листу CRM. Дубли с уже записанными в CRM
заказами (та же дата ±3 дня и тот же набор бутылок) не загружаются, поэтому импорт можно
повторять: второй прогон ничего не задвоит.
"""
import io
import re
from collections import Counter
from datetime import date, datetime, timedelta
from typing import Dict, List, Optional, Tuple

import openpyxl
from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from starlette.concurrency import run_in_threadpool

import crm
from auth import require_admin
from OrderItem import OrderItem
from UseDatabase import UseDatabase, config

# Колонки масел листа «Заказы» → названия как в прайсе CRM.
OIL_COLUMNS = {
    'подс': 'подсолнечное', 'рыж': 'рыжиковое', 'горч': 'горчичное', 'лен': 'льняное',
    'кокос': 'кокосовое', 'кун': 'кунжут нешлифованный', 'чкун': 'кунжут черный',
    'коноп': 'конопляное', 'раст': 'расторопши', 'грец': 'грецкого ореха',
    'тык': 'тыквы штирийской', 'минд': 'миндальное', 'кедр': 'кедровое', 'тмин': 'черного тмина',
    'облеп': 'облепиховое', 'амар': 'амарантовое',
}
SOURCES = (('авито', 'Avito'), ('avito', 'Avito'), ('озон', 'Ozon'), ('ozon', 'Ozon'), ('яндекс', 'Яндекс'))
RECENT_DAYS_OPEN = 7  # заказы с выдачей не раньше «неделю назад» считаем ещё открытыми


class OrderImportError(crm.CrmError):
    pass


def _norm(text) -> str:
    return re.sub(r'\s+', ' ', str(text or '').strip().lower().replace('ё', 'е'))


def _as_date(value) -> Optional[date]:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        match = re.fullmatch(r'\s*(\d{1,2})\.(\d{1,2})\.(\d{4})\s*', value)
        if match:
            day, month, year = map(int, match.groups())
            for shift in range(4):  # «31.09.2026» → последний реальный день месяца
                try:
                    return date(year, month, day - shift)
                except ValueError:
                    continue
    return None


def _as_ml(value) -> int:
    if value is None or value == '':
        return 0
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return int(round(value))
    try:
        return crm.parse_volume(value)
    except crm.CrmError:
        return 0


def read_price_sheet(workbook) -> Dict[str, Dict[int, float]]:
    """Лист «Прайс»: масло → {объём мл: цена}."""
    if 'Прайс' not in workbook.sheetnames:
        return {}
    sheet = workbook['Прайс']
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return {}
    volumes = {}
    for index, title in enumerate(rows[0]):
        match = re.fullmatch(r'\s*(\d+(?:[.,]\d+)?)\s*(мл|л)\s*', str(title or ''))
        if match:
            amount = float(match.group(1).replace(',', '.'))
            volumes[index] = int(amount * 1000 if match.group(2) == 'л' else amount)
    prices: Dict[str, Dict[int, float]] = {}
    for row in rows[1:]:
        name = _norm(row[0]) if row and row[0] else ''
        if not name or not volumes:
            if prices:
                break  # дальше на листе — урбеч и мука, это не масла
            continue
        for index, volume in volumes.items():
            value = row[index] if index < len(row) else None
            if isinstance(value, (int, float)) and value > 0:
                prices.setdefault(name, {})[volume] = round(float(value), 2)
    return prices


def split_into_bottles(volume_ml: int, sizes: List[int]) -> Optional[List[int]]:
    """Точное разложение объёма на бутылки прайса (меньше бутылок — лучше), иначе None."""
    sizes = sorted({size for size in sizes if size > 0}, reverse=True)
    if not sizes or volume_ml <= 0:
        return None
    step = 50
    if volume_ml % step:
        return None
    target = volume_ml // step
    best: List[Optional[List[int]]] = [None] * (target + 1)
    best[0] = []
    for amount in range(1, target + 1):
        for size in sizes:
            units = size // step
            if size % step == 0 and units <= amount and best[amount - units] is not None:
                candidate = best[amount - units] + [size]
                if best[amount] is None or len(candidate) < len(best[amount]):
                    best[amount] = candidate
    return best[target]


def price_line(oil: str, volume_ml: int, prices: Dict[int, float]) -> Tuple[List[dict], str]:
    """Позиции заказа для одной ячейки: [(объём, кол-во, цена)] и пометка ('' | 'approx' | 'unpriced')."""
    if not prices:
        return [{'volume': volume_ml, 'count': 1, 'price': 0.0}], 'unpriced'
    bottles = split_into_bottles(volume_ml, list(prices))
    if bottles:
        counts = Counter(bottles)
        return [{'volume': size, 'count': count, 'price': prices[size]} for size, count in sorted(counts.items())], ''
    # Нестандартный объём (170, 380 мл…): одна позиция по цене за мл ближайшей бутылки.
    nearest = min(prices, key=lambda size: (abs(size - volume_ml), -size))
    return [{'volume': volume_ml, 'count': 1, 'price': round(prices[nearest] / nearest * volume_ml)}], 'approx'


def parse_workbook(content: bytes, today: Optional[date] = None) -> dict:
    """Разобрать книгу в список заказов (без записи в базу)."""
    today = today or date.today()
    try:
        values_book = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
        formulas_book = openpyxl.load_workbook(io.BytesIO(content), data_only=False)
    except Exception as error:  # noqa: BLE001 — битый или не тот файл
        raise OrderImportError(f'Не удалось открыть файл как Excel (.xlsx): {error}')
    if 'Заказы' not in values_book.sheetnames:
        raise OrderImportError('В книге нет листа «Заказы»')

    sheet, formulas = values_book['Заказы'], formulas_book['Заказы']
    header = [_norm(value) for value in next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))]
    columns = {name: index for index, name in enumerate(header) if name}
    oil_columns = {index: OIL_COLUMNS[name] for name, index in columns.items() if name in OIL_COLUMNS}
    if len(oil_columns) < 5 or 'фио' not in columns:
        raise OrderImportError('Не узнал колонки: нужны масла (подс, лен, тык…) и «ФИО» в первой строке')
    col = {key: columns.get(key) for key in ('выдача', 'доп', 'фио', 'доставка', 'комментарии', 'тел', 'откуда')}
    order_date_col = (col['выдача'] or 2) - 1

    sheet_prices = read_price_sheet(values_book)
    orders, warnings, skipped_rows = [], [], 0
    current = None
    for row_number, (row, formula_row) in enumerate(
        zip(sheet.iter_rows(min_row=2, values_only=True), formulas.iter_rows(min_row=2, values_only=True)), start=2
    ):
        if any(isinstance(value, str) and value.startswith('=') for value in formula_row):
            skipped_rows += 1
            current = None  # строка итогов закрывает блок
            continue

        def cell(key):
            index = col.get(key) if key in col else None
            return row[index] if index is not None and index < len(row) else None

        bottles = [(oil, _as_ml(row[index])) for index, oil in oil_columns.items()
                   if index < len(row) and _as_ml(row[index]) > 0]
        order_day = _as_date(row[order_date_col] if order_date_col < len(row) else None)
        issue_day = _as_date(cell('выдача'))
        name = str(cell('фио') or '').strip()
        label = row[col['выдача']] if col['выдача'] is not None and col['выдача'] < len(row) else None
        if isinstance(label, str) and not issue_day and not name:
            current = None  # нижний блок итогов: «общая», «шт», «Прибыль»…
            continue

        if not order_day and not issue_day and not name:
            if bottles and current is not None:
                current['bottles'].extend(bottles)
                current['rows'].append(row_number)
            elif bottles:
                warnings.append(f'Стр. {row_number}: позиции без заказа — пропущены')
            continue
        if not bottles:
            if str(cell('доп') or '').strip():
                warnings.append(f'Стр. {row_number}: только допы без масел ({name or "без имени"}) — пропущено')
            current = None
            continue

        day = order_day or issue_day
        if not day and orders:
            day = orders[-1]['order_date']
            warnings.append(f'Стр. {row_number}: нет даты у «{name}» — взял дату предыдущего заказа ({day})')
        if day and day.year < 2020 and issue_day:
            warnings.append(f'Стр. {row_number}: дата заказа {day} похожа на опечатку — взял дату выдачи')
            day = issue_day
        if not day:
            warnings.append(f'Стр. {row_number}: нет даты — пропущено')
            current = None
            continue
        shipping = issue_day if issue_day and -60 <= (issue_day - day).days <= 150 else day

        delivery = str(cell('доставка') or '').strip()
        comment = str(cell('комментарии') or '').strip()
        extra = str(cell('доп') or '').strip()
        phone_raw = str(cell('тел') or '')
        phone_match = re.search(r'(?:\+7|8|7)?[\s(-]*\d{3}[\s)-]*\d{3}[\s-]*\d{2}[\s-]*\d{2}', phone_raw + ' ' + comment)
        source_text = _norm(f"{delivery} {cell('откуда') or ''}")
        source = next((title for key, title in SOURCES if key in source_text), None)
        current = {
            'row': row_number,
            'rows': [row_number],
            'order_date': day,
            'shipping_date': shipping,
            'customer_name': name or 'Без имени',
            'phone': phone_match.group(0).strip() if phone_match else '',
            'source': source,
            'note_parts': [part for part in (
                f'доставка: {delivery}' if delivery else '',
                comment,
                f'доп: {extra}' if extra and extra not in ('+',) else ('есть допы (урбеч/жмых/мука)' if extra else ''),
                f"откуда: {cell('откуда')}" if cell('откуда') else '',
            ) if part],
            'bottles': bottles,
        }
        orders.append(current)

    crm_prices = _crm_price_index()
    month_totals: Dict[str, dict] = {}
    approx = unpriced = 0
    for order in orders:
        items: Dict[Tuple[str, int, float], int] = {}
        flags = set()
        for oil, volume in order['bottles']:
            prices = sheet_prices.get(_norm(oil)) or crm_prices.get(_norm(oil), {})
            lines, flag = price_line(oil, volume, prices)
            if flag:
                flags.add(flag)
            for line in lines:
                key = (oil, line['volume'], line['price'])
                items[key] = items.get(key, 0) + line['count']
        order['items'] = [{'oil_name': oil, 'volume': volume, 'price': price, 'count': count}
                          for (oil, volume, price), count in items.items()]
        order['total'] = sum(item['price'] * item['count'] for item in order['items'])
        order['approx'] = 'approx' in flags
        order['unpriced'] = 'unpriced' in flags
        approx += order['approx']
        unpriced += order['unpriced']
        open_order = order['shipping_date'] >= today - timedelta(days=RECENT_DAYS_OPEN)
        order['is_open'] = open_order
        order['note'] = '; '.join(order['note_parts'] + [f"импорт из Excel, стр. {order['row']}"])[:255]
        month = order['order_date'].strftime('%Y-%m')
        bucket = month_totals.setdefault(month, {'month': month, 'orders': 0, 'revenue': 0.0, 'bottles': 0})
        bucket['orders'] += 1
        bucket['revenue'] += order['total']
        bucket['bottles'] += sum(item['count'] for item in order['items'])

    return {
        'orders': orders,
        'months': [month_totals[key] for key in sorted(month_totals)],
        'skipped_rows': skipped_rows,
        'approx_orders': approx,
        'unpriced_orders': unpriced,
        'warnings': warnings,
        'price_source': 'лист «Прайс» из файла' if sheet_prices else 'прайс-лист CRM',
    }


def _crm_price_index() -> Dict[str, Dict[int, float]]:
    with UseDatabase(config) as cursor:
        cursor.execute('SELECT oil_name, volume, price FROM price_list WHERE price > 0 ORDER BY id')
        index: Dict[str, Dict[int, float]] = {}
        for name, volume, price in cursor.fetchall():
            index.setdefault(_norm(name), {}).setdefault(int(volume), float(price))
    return index


def _item_signature(items) -> Tuple:
    return tuple(sorted((_norm(item['oil_name']), int(item['volume']), int(item['count'])) for item in items))


def find_duplicates(orders: List[dict]) -> set:
    """Номера строк Excel, чьи заказы уже есть в CRM (дата ±3 дня и тот же набор бутылок)."""
    if not orders:
        return set()
    start = min(order['order_date'] for order in orders) - timedelta(days=3)
    end = max(order['order_date'] for order in orders) + timedelta(days=3)
    with UseDatabase(config) as cursor:
        cursor.execute(
            """SELECT o.id, o.date, d.oil_name, d.volume, d.count
               FROM orders o JOIN order_details d ON d.order_id = o.id
               WHERE o.date BETWEEN %s AND %s""",
            (start, end),
        )
        existing: Dict[int, dict] = {}
        for order_id, order_day, oil, volume, count in cursor.fetchall():
            entry = existing.setdefault(order_id, {'date': order_day, 'items': []})
            entry['items'].append({'oil_name': oil, 'volume': volume, 'count': count})
    by_signature: Dict[Tuple, List[date]] = {}
    for entry in existing.values():
        day = entry['date'].date() if isinstance(entry['date'], datetime) else entry['date']
        by_signature.setdefault(_item_signature(entry['items']), []).append(day)
    duplicates = set()
    for order in orders:
        days = by_signature.get(_item_signature(order['items']), [])
        if any(abs((day - order['order_date']).days) <= 3 for day in days):
            duplicates.add(order['row'])
    return duplicates


def run_import(content: bytes, commit: bool = False, today: Optional[date] = None) -> dict:
    """Проверка (commit=False) или загрузка заказов из книги. Возвращает отчёт для экрана."""
    crm.ensure_crm_schema()
    parsed = parse_workbook(content, today)
    orders = parsed['orders']
    duplicates = find_duplicates(orders)
    fresh = [order for order in orders if order['row'] not in duplicates]

    existing_customers = {_norm(row['full_name']): row['id'] for row in crm.list_customers(limit=100000)}
    new_customer_names = {_norm(order['customer_name']) for order in fresh} - set(existing_customers)

    report = {
        'orders_in_file': len(orders),
        'duplicates': len(duplicates),
        'to_import': len(fresh),
        'new_customers': len(new_customer_names),
        'open_orders': sum(order['is_open'] for order in fresh),
        'approx_orders': parsed['approx_orders'],
        'unpriced_orders': parsed['unpriced_orders'],
        'skipped_rows': parsed['skipped_rows'],
        'price_source': parsed['price_source'],
        'months': parsed['months'],
        'revenue_total': sum(month['revenue'] for month in parsed['months']),
        'warnings': parsed['warnings'][:50],
        'sample': [_preview(order, order['row'] in duplicates) for order in orders[:15]],
        'imported': 0,
    }
    if commit and fresh:
        report['imported'] = _write_orders(fresh, existing_customers)
    return report


def _preview(order: dict, duplicate: bool) -> dict:
    return {
        'row': order['row'],
        'date': order['order_date'].isoformat(),
        'customer': order['customer_name'],
        'items': ', '.join(f"{item['oil_name']} {item['volume']} мл × {item['count']}" for item in order['items']),
        'total': order['total'],
        'approx': order['approx'],
        'duplicate': duplicate,
    }


def _write_orders(orders: List[dict], customers: Dict[str, int]) -> int:
    """Записать заказы одной транзакцией: ошибка на середине не оставит половину импорта."""
    done_status = (crm.status_ids('done') or [crm.default_status_id()])[0]
    open_status = crm.default_status_id()
    statuses = {row['id']: row['name'] for row in crm.list_statuses()}
    seed_info = {}
    with UseDatabase(config) as cursor:
        cursor.execute('SELECT oil_name, volume, seed_weight_kg, seed_price_per_kg FROM price_list')
        for name, volume, seed_kg, seed_price in cursor.fetchall():
            seed_info.setdefault((_norm(name), int(volume)), (float(seed_kg or 0), float(seed_price or 0)))

    with UseDatabase(config) as cursor:
        for order in orders:
            key = _norm(order['customer_name'])
            customer_id = customers.get(key)
            if customer_id is None:
                first, _, rest = order['customer_name'].partition(' ')
                cursor.execute(
                    """INSERT INTO customers (name, surname, phone, address, source, note, created_at)
                       VALUES (%s, %s, %s, '', %s, %s, %s)""",
                    (first[:50] or 'Без имени', rest[:100], order['phone'][:32], order['source'],
                     'из старой таблицы заказов', order['order_date']),
                )
                customer_id = cursor.lastrowid
                customers[key] = customer_id
            products = []
            for item in order['items']:
                seed_kg, seed_price = seed_info.get((_norm(item['oil_name']), int(item['volume'])), (0.0, 0.0))
                products.append(OrderItem(item['oil_name'], int(item['volume']), int(item['count']),
                                          float(item['price']), seed_kg, seed_price))
            status = open_status if order['is_open'] else done_status
            cursor.execute(
                """INSERT INTO orders (customer_id, date, shipping_date, total_price, status, note)
                   VALUES (%s, %s, %s, %s, %s, %s)""",
                (customer_id, order['order_date'], order['shipping_date'],
                 sum(product.total_revenue() for product in products), status, order['note']),
            )
            order_id = cursor.lastrowid
            for product in products:
                cursor.execute(
                    """INSERT INTO order_details (
                            order_id, oil_name, volume, count, price,
                            seed_weight_kg, seed_price_per_kg, seed_cost, revenue, profit
                       ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s)""",
                    (order_id, product.oil_name, product.volume, product.count, product.price,
                     product.seed_weight_kg, product.seed_price_per_kg, product.total_seed_cost(),
                     product.total_revenue(), product.total_profit()),
                )
            cursor.execute(
                "INSERT INTO statuses_list (order_id, order_status, name, datetime) VALUES (%s, %s, %s, %s)",
                (order_id, status, statuses.get(status, ''), datetime.combine(order['order_date'], datetime.min.time())),
            )
    return len(orders)


# ── HTTP ─────────────────────────────────────────────────────────────────────

MAX_FILE_BYTES = 10 * 1024 * 1024
router = APIRouter(prefix='/api/crm/import')


@router.post('/orders')
async def api_import_orders(
    file: UploadFile = File(...),
    commit: bool = Form(False),
    admin: dict = Depends(require_admin),
):
    """Сначала проверка (commit=false) — отчёт и выручка по месяцам, потом загрузка (commit=true)."""
    content = await file.read()
    if len(content) > MAX_FILE_BYTES:
        raise HTTPException(status_code=400, detail='Файл больше 10 МБ')
    if not (file.filename or '').lower().endswith(('.xlsx', '.xlsm')):
        raise HTTPException(status_code=400, detail='Нужен файл Excel .xlsx')
    try:
        return await run_in_threadpool(run_import, content, commit)
    except crm.CrmError as error:
        raise HTTPException(status_code=400, detail=str(error))
