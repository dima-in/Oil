"""Разбор старой таблицы заказов: продолжения строк, итоги, разложение объёма на бутылки."""
import io
from datetime import date, datetime

import openpyxl
import pytest

import order_import

HEADER = [None, None, 'выдача', 'подс', 'рыж', 'горч', 'лен', 'кокос', 'кун', 'чкун', 'коноп', 'раст', 'грец',
          'тык', 'минд', 'кедр', 'тмин', 'облеп', 'амар', 'доп', 'ФИО', 'доставка', 'комментарии', 'тел', 'откуда']
COL = {name: index for index, name in enumerate(HEADER) if name}


def order_row(day=None, issue=None, name=None, delivery=None, comment=None, **oils):
    row = [None] * len(HEADER)
    row[1], row[2] = day, issue
    row[COL['ФИО']], row[COL['доставка']], row[COL['комментарии']] = name, delivery, comment
    for oil, volume in oils.items():
        row[COL[oil]] = volume
    return row


def build_workbook(rows):
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = 'Заказы'
    sheet.append(HEADER)
    sheet.append([])
    for row in rows:
        sheet.append(row)
    prices = book.create_sheet('Прайс')
    prices.append(['масло', '50 мл', '100 мл', '200 мл', '250 мл', '500 мл', '1 л'])
    prices.append(['льняное', None, None, None, 870, 1560, 2800])
    prices.append(['конопляное', None, 650, None, 1290, 2320, 4170])
    prices.append(['облепиховое', 1970, 3430, 6170])
    prices.append([])
    prices.append(['урбеч', '200 гр'])  # ниже масел — не прайс масел
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


@pytest.fixture(autouse=True)
def no_db(monkeypatch):
    monkeypatch.setattr(order_import, '_crm_price_index', lambda: {})


def test_split_into_bottles():
    assert order_import.split_into_bottles(2000, [250, 500, 1000]) == [1000, 1000]
    assert sorted(order_import.split_into_bottles(750, [250, 500, 1000])) == [250, 500]
    assert order_import.split_into_bottles(300, [100, 250, 500]) == [100, 100, 100]
    assert order_import.split_into_bottles(170, [100, 250]) is None


def test_price_line_marks_approximate_and_unpriced():
    lines, flag = order_import.price_line('кокосовое', 380, {200: 810})
    assert flag == 'approx' and lines == [{'volume': 380, 'count': 1, 'price': 1539}]
    lines, flag = order_import.price_line('грецкого ореха', 250, {})
    assert flag == 'unpriced' and lines[0]['price'] == 0.0


def test_parse_workbook_groups_rows_and_skips_totals():
    rows = [
        order_row(datetime(2025, 3, 2), datetime(2025, 3, 6), 'Danny Galimov', 'св', коноп=2000),
        order_row(datetime(2025, 3, 4), '31.03.2025', 'Анна', 'авито', 'ждет', лен=500, облеп=400),
        order_row(**{'лен': 250}),  # продолжение заказа Анны
        order_row(None, None, 'Без даты'),  # без масел — пропускается
        ['=SUM(D3:D5)'] + [None] * (len(HEADER) - 1),  # строка итогов
        order_row(datetime(2015, 3, 9), datetime(2025, 3, 12), 'Опечатка', 'озон', лен=250),
    ]
    parsed = order_import.parse_workbook(build_workbook(rows), today=date(2026, 10, 22))
    orders = parsed['orders']
    assert [order['customer_name'] for order in orders] == ['Danny Galimov', 'Анна', 'Опечатка']

    danny, anna, typo = orders
    assert danny['items'] == [{'oil_name': 'конопляное', 'volume': 1000, 'price': 4170.0, 'count': 2}]
    assert danny['total'] == 8340
    # 500 + 250 льняного (строка-продолжение) и 2 × 200 облепихи
    assert sorted((item['oil_name'], item['volume'], item['count']) for item in anna['items']) == [
        ('льняное', 250, 1), ('льняное', 500, 1), ('облепиховое', 200, 2)]
    assert anna['total'] == 870 + 1560 + 2 * 6170
    assert anna['shipping_date'] == date(2025, 3, 31)
    assert anna['source'] == 'Avito' and 'ждет' in anna['note']
    assert typo['order_date'] == date(2025, 3, 12)  # 2015 → дата выдачи
    assert typo['source'] == 'Ozon'
    assert parsed['skipped_rows'] == 1
    assert [month['month'] for month in parsed['months']] == ['2025-03']
    assert parsed['months'][0]['revenue'] == danny['total'] + anna['total'] + typo['total']
    assert not any(order['is_open'] for order in orders)


def test_parse_workbook_rejects_wrong_file():
    with pytest.raises(order_import.OrderImportError):
        order_import.parse_workbook(b'not an excel file')
    book = openpyxl.Workbook()
    book.active.title = 'Лист1'
    buffer = io.BytesIO()
    book.save(buffer)
    with pytest.raises(order_import.OrderImportError, match='Заказы'):
        order_import.parse_workbook(buffer.getvalue())


def economics_workbook():
    book = openpyxl.Workbook()
    price = book.active
    price.title = 'Прайс'
    price.append(['масло', '250 мл', '500 мл', 'доставка за 1 кг', 'бутылка', 'кг в закладке', 'цена за 1 кг',
                  'получка за закладку', 'выход масла с закладки'])
    price.append(['подсолнечное', '=ROUNDDOWN(((H2+G2*F2)/I2*225+D2*F2+E2)*1.06, -1)', None, 50, 23, 2, 149, 2300,
                  1000])
    price.append(['льняное', None, None, None, None, 1.8, 105, 2200, 750])
    price.append(['кунжут черный'])  # без норм — пропускается
    pnl = book.create_sheet('Прибыль и убыль')
    pnl.append([None, 'дата', 'наименование', 'кг', 'цена за 1 кг', 'сумма заказа', 'доставка', 'каршеринг'])
    pnl.append([None, datetime(2026, 3, 3), 'тмин', 25, 350, 8750])
    pnl.append([None, None, 'кунжут', 25, 180, 4500, None, 1700])  # дата — как строкой выше
    pnl.append([None, datetime(2026, 1, 26), 'Конопля 40 кг', None, None, 0])  # без суммы
    pnl.append([None, None, None, None, None, 50000])  # итог без наименования
    values_buffer = io.BytesIO()
    book.save(values_buffer)
    return values_buffer.getvalue()


def test_read_economics_and_expenses():
    content = economics_workbook()
    values = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    formulas = openpyxl.load_workbook(io.BytesIO(content), data_only=False)

    economics = order_import.read_economics(values, formulas)
    assert economics['settings'] == {'seed_delivery_per_kg': 50.0, 'bottle_cost': 23.0, 'tax_percent': 6.0}
    sunflower, flax = economics['oils']
    assert sunflower == {'oil_name': 'подсолнечное', 'kg_per_batch': 2.0, 'yield_ml': 1000.0, 'yield_percent': 50.0,
                         'seed_price_per_kg': 149.0, 'labor_per_batch': 2300.0}
    assert flax['yield_percent'] == round(750 / 18, 4)

    expenses, warnings = order_import.read_expenses(values)
    assert [(e['expense_date'], e['item_name'], e['goods_total'], e['carsharing_cost']) for e in expenses] == [
        (date(2026, 3, 3), 'тмин', 8750.0, 0.0),
        (date(2026, 3, 3), 'кунжут', 4500.0, 1700.0),
    ]
    assert any('Конопля 40 кг' in warning for warning in warnings)
